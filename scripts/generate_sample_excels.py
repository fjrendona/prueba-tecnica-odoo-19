"""Genera los Excel de ejemplo de samples/.

- ventas_junio_2026.xlsx: ventas válidas de junio de 2026. Incluye a propósito
  variantes que el importador debe tolerar: mayúsculas y espacios en los nombres,
  estados escritos de distintas formas, números como texto en formato colombiano,
  fechas como texto, una cantidad decimal y una columna extra que se ignora.
- ventas_julio_2026.xlsx: ventas válidas de julio de 2026. Mezcla clientes,
  productos y vendedores ya creados en junio (se reutilizan, escritos con otras
  mayúsculas) con otros nuevos, para validar el comportamiento acumulado: el
  tablero muestra dos meses y el lote de julio solo cuenta los maestros nuevos.
- ventas_agosto_2026.xlsx: ventas válidas de agosto de 2026, pensadas para
  importarlas con la lista de ventas abierta en otra sesión y ver la
  actualización en tiempo real.
- ventas_con_errores.xlsx: un error de cada tipo, para ver el reporte completo.

Uso (desde la raíz del repo): python3 scripts/generate_sample_excels.py
"""

import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "samples"
HEADERS = ["Fecha", "Cliente", "Vendedor", "Producto", "Cantidad", "Valor Unitario", "Valor Total", "Estado"]

# (día, cliente, vendedor, producto, cantidad, valor unitario, estado)
JUNE_SALES = [
    (1, "Distribuidora El Sol", "Juan Pérez", "Panel Solar 450W", 2, 850_000, "Confirmada"),
    (2, "Supermercado La 14", "María Gómez", "Inversor 5kW", 1, 3_200_000, "Confirmada"),
    (3, "Ferretería Central", "Carlos Ruiz", "Batería Litio 5kWh", 1, 6_500_000, "Borrador"),
    (4, "distribuidora el sol ", "Juan Pérez", "Cable Solar 6mm (metro)", 120, 4_500, "confirmado"),
    (5, "Hotel Andino", "María Gómez", "Panel Solar 450W", 12, 820_000, "Confirmada"),
    (6, "Cafetería Aroma", "Carlos Ruiz", "Controlador MPPT 60A", 2, 780_000, "CONFIRMADA"),
    (8, "Supermercado La 14", "Juan Pérez", "Estructura de Montaje", 6, 350_000, "Confirmada"),
    (9, "Ferretería Central", "Carlos Ruiz", "Cable Solar 6mm (metro)", 250, 4_200, "Confirmada"),
    (10, "Hotel Andino", "María Gómez", "Inversor 5kW", 2, 3_150_000, "Cancelada"),
    (11, "Colegio San José", "Ana Torres", "Panel Solar 450W", 20, 800_000, "Confirmada"),
    (12, "Cafetería Aroma", "Ana Torres", "Batería Litio 5kWh", 1, 6_400_000, "Borrador"),
    (15, "HOTEL ANDINO", "juan pérez", "Estructura de Montaje", 10, 340_000, "Confirmada"),
    (16, "Colegio San José", "Ana Torres", "Controlador MPPT 60A", 4, 760_000, "Confirmada"),
    (17, "Distribuidora El Sol", "María Gómez", "Cable Solar 6mm (metro)", 80, 4_500, "Anulada"),
    (18, "Ferretería Central", "Carlos Ruiz", "Panel Solar 450W", 4, 845_000, "Confirmada"),
    (19, "Supermercado La 14", "Juan Pérez", "Batería Litio 5kWh", 2, 6_300_000, "Confirmada"),
    (22, "Clínica Norte", "Ana Torres", "Inversor 5kW", 3, 3_100_000, "Confirmada"),
    (23, "Clínica Norte", "Ana Torres", "Panel Solar 450W", 30, 790_000, "Borrador"),
    (24, "Cafetería Aroma", "Carlos Ruiz", "Estructura de Montaje", 2, 355_000, "Confirmada"),
    (25, "Colegio San José", "María Gómez", "Cable Solar 6mm (metro)", 150, 4_300, "Confirmada"),
    (26, "Hotel Andino", "Juan Pérez", "Controlador MPPT 60A", 3, 770_000, "Confirmada"),
    (29, "Distribuidora El Sol", "María Gómez", "Inversor 5kW", 1, 3_250_000, "Confirmada"),
    (30, "Clínica Norte", "Carlos Ruiz", "Batería Litio 5kWh", 1, 6_450_000, "Confirmada"),
]
# Fila extra de junio: cantidad decimal escrita como texto con coma (12,5 × 4.300).
JUNE_EXTRA_ROWS = [
    [datetime.date(2026, 6, 30), "Colegio San José", "Ana Torres", "Cable Solar 6mm (metro)",
     "12,5", 4300, 53750, "Confirmada", "Cantidad decimal como texto"],
]
JUNE_EXTRA_CONFIRMED = (1, 53_750)  # (ventas, total) de JUNE_EXTRA_ROWS

JULY_SALES = [
    # Maestros de junio, reutilizados (con otras mayúsculas y espacios)
    (1, "Distribuidora El Sol", "Juan Pérez", "Panel Solar 450W", 10, 840_000, "Confirmada"),
    (2, "HOTEL ANDINO", "maría gómez", "Inversor 5kW", 2, 3_150_000, "Confirmada"),
    (3, "Supermercado La 14", "Carlos Ruiz", "Batería Litio 5kWh", 3, 6_300_000, "Borrador"),
    (6, "colegio san josé", "Ana Torres", "Estructura de Montaje", 8, 345_000, "Confirmada"),
    (7, "Clínica Norte", "Juan Pérez", "Cable Solar 6mm (metro)", 300, 4_200, "Confirmado"),
    (8, "Cafetería Aroma", "María Gómez", "Panel Solar 450W", 4, 830_000, "Cancelada"),
    # Clientes, producto y vendedora nuevos
    (9, "Panadería La Espiga", "Laura Méndez", "Microinversor 800W", 6, 1_150_000, "Confirmada"),
    (10, "Constructora Horizonte", "Laura Méndez", "Panel Solar 450W", 40, 790_000, "Confirmada"),
    (13, "Constructora Horizonte", "Carlos Ruiz", "Inversor 5kW", 5, 3_050_000, "Borrador"),
    (14, "Panadería La Espiga", "Laura Méndez", "Batería Litio 5kWh", 1, 6_350_000, "Confirmada"),
    (15, "Ferretería Central", "Juan Pérez", "Microinversor 800W", 12, 1_120_000, "Confirmada"),
    (16, "Hotel Andino", "Ana Torres", "Controlador MPPT 60A", 6, 765_000, "Anulada"),
    (20, "Distribuidora El Sol", "Laura Méndez", "Cable Solar 6mm (metro)", 200, 4_400, "Confirmada"),
    (21, "Constructora Horizonte", "Laura Méndez", "Estructura de Montaje", 20, 330_000, "Confirmada"),
    (22, "Clínica Norte", "María Gómez", "Microinversor 800W", 8, 1_140_000, "Borrador"),
    (27, "Supermercado La 14", "Juan Pérez", "Panel Solar 450W", 15, 815_000, "Confirmada"),
    (28, "Colegio San José", "Ana Torres", "Controlador MPPT 60A", 3, 770_000, "Confirmada"),
    (31, "Panadería La Espiga", "Carlos Ruiz", "Panel Solar 450W", 6, 835_000, "CONFIRMADA"),
]

AUGUST_SALES = [
    (3, "Constructora Horizonte", "Laura Méndez", "Panel Solar 450W", 25, 780_000, "Confirmada"),
    (4, "Hotel Andino", "María Gómez", "Batería Litio 5kWh", 2, 6_250_000, "Confirmada"),
    (5, "Panadería La Espiga", "Carlos Ruiz", "Microinversor 800W", 4, 1_130_000, "Borrador"),
    (6, "Distribuidora El Sol", "Juan Pérez", "Inversor 5kW", 3, 3_100_000, "Confirmada"),
    (10, "Colegio San José", "Ana Torres", "Cable Solar 6mm (metro)", 180, 4_350, "Confirmada"),
    (11, "Clínica Norte", "Laura Méndez", "Controlador MPPT 60A", 5, 760_000, "Confirmada"),
    (12, "Supermercado La 14", "Juan Pérez", "Estructura de Montaje", 12, 340_000, "Cancelada"),
    (13, "Ferretería Central", "Carlos Ruiz", "Panel Solar 450W", 8, 825_000, "Confirmada"),
    (17, "Cafetería Aroma", "María Gómez", "Microinversor 800W", 2, 1_150_000, "Confirmada"),
    (18, "Constructora Horizonte", "Laura Méndez", "Inversor 5kW", 4, 3_050_000, "Borrador"),
    (19, "Hotel Andino", "Ana Torres", "Panel Solar 450W", 18, 800_000, "Confirmada"),
    (20, "Distribuidora El Sol", "Juan Pérez", "Batería Litio 5kWh", 1, 6_400_000, "Confirmada"),
    (24, "Panadería La Espiga", "Laura Méndez", "Estructura de Montaje", 6, 350_000, "Confirmada"),
    (25, "Clínica Norte", "Carlos Ruiz", "Panel Solar 450W", 22, 795_000, "Confirmada"),
    (26, "Colegio San José", "María Gómez", "Inversor 5kW", 2, 3_120_000, "Anulada"),
    (28, "Supermercado La 14", "Juan Pérez", "Controlador MPPT 60A", 4, 770_000, "Confirmada"),
]


def colombian(number):
    """1234.5 -> "1.234,5" (texto como lo escribiría una persona en Colombia)."""
    integer, _, decimals = f"{number:.2f}".partition(".")
    text = f"{int(integer):,}".replace(",", ".")
    decimals = decimals.rstrip("0")
    return f"{text},{decimals}" if decimals else text


def new_sheet(headers, title):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="714B67")
    for column, width in zip("ABCDEFGHI", (12, 26, 16, 28, 10, 16, 16, 12, 30)):
        sheet.column_dimensions[column].width = width
    return workbook, sheet


def build_sales_workbook(title, month, sales, extra_rows=()):
    workbook, sheet = new_sheet([*HEADERS, "Observaciones"], title)
    for index, (day, partner, salesperson, product, quantity, price_unit, state) in enumerate(sales):
        date = datetime.date(2026, month, day)
        total = quantity * price_unit
        row = [date, partner, salesperson, product, quantity, price_unit, total, state, ""]
        if index % 5 == 1:  # algunas fechas como texto
            row[0] = date.strftime("%d/%m/%Y")
        if index % 4 == 2:  # algunos importes como texto en formato colombiano
            row[5], row[6] = colombian(price_unit), colombian(total)
            row[8] = "Importes escritos como texto"
        sheet.append(row)
        if isinstance(row[0], datetime.date):
            sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd"
    for row in extra_rows:
        sheet.append(row)
        sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm-dd"
    return workbook


def build_invalid_workbook():
    workbook, sheet = new_sheet(HEADERS, "Ventas con errores")
    rows = [
        ["2026-06-01", "Distribuidora El Sol", "Juan Pérez", "Panel Solar 450W", 2, 850000, 1700000, "Confirmada"],
        ["2026-06-02", None, "María Gómez", "Inversor 5kW", 1, 3200000, 3200000, "Confirmada"],
        ["2026-06-03", "Ferretería Central", "Carlos Ruiz", "Batería Litio 5kWh", 0, 6500000, 6500000, "Borrador"],
        ["2026-06-04", "Hotel Andino", "María Gómez", "Panel Solar 450W", 12, 820000, 9000000, "Confirmada"],
        ["2026-06-05", "Cafetería Aroma", "Carlos Ruiz", "Controlador MPPT 60A", 2, 780000, 1560000, "Pendiente"],
        ["31/02/2026", "Supermercado La 14", "Juan Pérez", "Estructura de Montaje", 6, 350000, 2100000, "Confirmada"],
        ["2026-06-07", "Colegio San José", "Ana Torres", "Panel Solar 450W", "dos", 800000, 1600000, "Confirmada"],
        ["2026-06-08", "Clínica Norte", None, "Inversor 5kW", 3, 3100000, 9300000, "Confirmada"],
        ["2026-06-09", "Clínica Norte", "Ana Torres", "Inversor 5kW", 1, 0, 0, "Confirmada"],
        [None, None, None, None, None, None, None, None],  # fila vacía: se ignora
        ["2026-06-10", "Hotel Andino", "Juan Pérez", None, 3, 770000, 2310000, None],
    ]
    for row in rows:
        sheet.append(row)
    return workbook


def summarize(file_name, sales, extra_count=0, extra_confirmed=(0, 0)):
    """Imprime las cifras esperadas en el tablero para validar la importación."""
    by_state = {}
    for sale in sales:
        state = sale[6].lower()
        key = "confirmada" if state.startswith("confirm") else "borrador" if state == "borrador" else "cancelada"
        count, total = by_state.get(key, (0, 0))
        by_state[key] = (count + 1, total + sale[4] * sale[5])
    count, total = by_state.get("confirmada", (0, 0))
    by_state["confirmada"] = (count + extra_confirmed[0], total + extra_confirmed[1])
    detail = ", ".join(f"{state} {count} por {colombian(total)}" for state, (count, total) in sorted(by_state.items()))
    print(f"{file_name}: {len(sales) + extra_count} ventas ({detail})")


def main():
    SAMPLES_DIR.mkdir(exist_ok=True)
    build_sales_workbook("Ventas junio 2026", 6, JUNE_SALES, JUNE_EXTRA_ROWS).save(
        SAMPLES_DIR / "ventas_junio_2026.xlsx"
    )
    build_sales_workbook("Ventas julio 2026", 7, JULY_SALES).save(SAMPLES_DIR / "ventas_julio_2026.xlsx")
    build_sales_workbook("Ventas agosto 2026", 8, AUGUST_SALES).save(SAMPLES_DIR / "ventas_agosto_2026.xlsx")
    build_invalid_workbook().save(SAMPLES_DIR / "ventas_con_errores.xlsx")

    summarize("ventas_junio_2026.xlsx", JUNE_SALES, len(JUNE_EXTRA_ROWS), JUNE_EXTRA_CONFIRMED)
    summarize("ventas_julio_2026.xlsx", JULY_SALES)
    summarize("ventas_agosto_2026.xlsx", AUGUST_SALES)
    print("ventas_con_errores.xlsx: 11 filas (1 vacía), 10 errores esperados")


if __name__ == "__main__":
    main()
